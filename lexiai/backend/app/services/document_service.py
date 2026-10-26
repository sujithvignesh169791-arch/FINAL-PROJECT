import os
import pymupdf as fitz
from docx import Document as DocxDocument
import pytesseract
from PIL import Image
import json
import logging
from sqlalchemy.orm import Session
from app.db.session import SessionLocal
from app.models.document import Document
from app.models.analysis import DocumentAnalysis
import traceback

logger = logging.getLogger(__name__)

def extract_text_from_txt(file_path: str):
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
        return [{"page_number": 1, "text": text}], text, "txt"
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            text = f.read()
        return [{"page_number": 1, "text": text}], text, "txt"

def extract_text_from_docx(file_path: str):
    doc = DocxDocument(file_path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                lines.append(row_text)
    text = "\n".join(lines)
    return [{"page_number": 1, "text": text}], text, "docx"

def extract_text_from_pdf(file_path: str):
    doc = fitz.open(file_path)
    pages = []
    full_text = ""
    is_scanned = True

    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text = page.get_text("text").strip()
        
        if len(text) > 50:
            is_scanned = False
            
        pages.append({"page_number": page_num + 1, "text": text})
        full_text += text + "\n"
        
    method = "pymupdf"
    
    # If the document appears scanned (very little text extracted normally)
    if is_scanned:
        full_text = ""
        method = "ocr"
        ocr_succeeded = False
        for i, page_dict in enumerate(pages):
            page = doc.load_page(i)
            # Render page to an image
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            img_path = f"{file_path}_page_{i}.png"
            pix.save(img_path)
            
            try:
                # Perform OCR
                text = pytesseract.image_to_string(Image.open(img_path))
                page_dict["text"] = text
                full_text += text + "\n"
                ocr_succeeded = True
            except Exception as e:
                logger.warning(f"OCR failed or Tesseract is not installed: {e}")
                break
            finally:
                if os.path.exists(img_path):
                    try:
                        os.remove(img_path)
                    except Exception:
                        pass
                        
        if not ocr_succeeded:
            method = "pymupdf_fallback"
            full_text = "\n".join([p.get("text", "") for p in pages if p.get("text")])
            if not full_text.strip():
                full_text = "[Notice: This is a scanned or image-based document. Selectable text was not detected and OCR is not installed.]"
                pages = [{"page_number": 1, "text": full_text}]
                    
    return pages, full_text, method

def process_document(arg1, arg2=None):
    """
    Accepts either (document_id) or (db, document_id) for compatibility with tests
    and background tasks. Background tasks create their own DB session so they do not
    rely on request-scoped sessions that get closed by FastAPI.
    """
    own_session = False
    if isinstance(arg1, str) and arg2 is None:
        document_id = arg1
        db = SessionLocal()
        own_session = True
    elif arg2 is not None and isinstance(arg2, str):
        db = arg1
        document_id = arg2
    else:
        document_id = str(arg1)
        db = SessionLocal()
        own_session = True

    try:
        db_doc = db.query(Document).filter(Document.id == document_id).first()
        if not db_doc:
            return

        file_path = db_doc.storage_url
        ext = db_doc.file_type.lower()
        
        if ext == 'txt':
            pages, full_text, method = extract_text_from_txt(file_path)
        elif ext in ('docx', 'doc'):
            pages, full_text, method = extract_text_from_docx(file_path)
        elif ext == 'pdf':
            pages, full_text, method = extract_text_from_pdf(file_path)
        else:
            raise ValueError(f"Unsupported file extension: {ext}")
            
        # Create or update Analysis record
        analysis = db.query(DocumentAnalysis).filter(DocumentAnalysis.document_id == db_doc.id).first()
        if not analysis:
            analysis = DocumentAnalysis(document_id=db_doc.id)
            db.add(analysis)

        analysis.extracted_text = full_text
        analysis.page_text = pages
        analysis.processing_method = method
        analysis.processing_error = None
        
        # Update Document status
        db_doc.upload_status = "completed"
        db.commit()
    except Exception as e:
        db.rollback()
        logger.error(f"Error processing document {document_id}: {e}", exc_info=True)
        try:
            db_doc = db.query(Document).filter(Document.id == document_id).first()
            if db_doc:
                db_doc.upload_status = "failed"
                analysis = db.query(DocumentAnalysis).filter(DocumentAnalysis.document_id == db_doc.id).first()
                if not analysis:
                    analysis = DocumentAnalysis(document_id=db_doc.id)
                    db.add(analysis)
                analysis.processing_error = str(e)
                analysis.processing_method = "failed"
                db.commit()
        except Exception:
            pass
    finally:
        if own_session:
            db.close()

